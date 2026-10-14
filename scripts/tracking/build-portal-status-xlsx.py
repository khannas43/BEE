#!/usr/bin/env python3
"""Build docs/tracking/BEE-portal-status.xlsx from lib catalogue + work register."""
from __future__ import annotations

import json
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/tracking/BEE-portal-status.xlsx"
CATALOG = ROOT / "docs/tracking/.screen-catalog.json"
MATRIX_MD = ROOT / "docs/wp01/SCREEN_ACTION_MATRIX.md"


def disposition_to_handling(disposition: str) -> str:
    if disposition == "retire":
        return "retire"
    if disposition in ("merge", "contextual"):
        return "merge"
    return "keep"


def load_handling_by_app_path() -> dict[str, str]:
    text = MATRIX_MD.read_text(encoding="utf-8")
    out: dict[str, str] = {}
    for line in text.splitlines():
        if not line.startswith("| `/app/"):
            continue
        if " --- " in line:
            continue
        parts = [p.strip() for p in line.split("|")]
        if len(parts) < 4:
            continue
        path = parts[1].strip("`")
        match = __import__("re").match(r"^(retain|merge|contextual|retire|dev-only)", parts[3])
        if not match:
            continue
        out[path] = disposition_to_handling(match.group(1))
    return out


def apply_screen_handling(rows: list[dict], handling_by_path: dict[str, str]) -> None:
    for row in rows:
        sid = row["screenId"]
        handling = handling_by_path.get(f"/app/{sid}")
        if not handling and sid in ("administration/business-settings",):
            handling = "keep"
        if not handling and row.get("status") in ("Live – BEE service", "Planned"):
            handling = "keep"
        row["screenHandling"] = handling or ""

LEGEND = (
    "Legend — Live – BEE service: signed-in user; records and permissions come from the BEE database "
    "through the portal API (not preview/sample data). Prototype: catalogue UI with sample or preview data. "
    "Planned: named in scope, not yet live. Screen handling (WP01): keep = retained entry route; "
    "merge = folds into another route or is a tab/panel (contextual); retire = duplicate removed in favour of another route."
)

TAB2_ROWS = [
    ("Cursor brief", "BRIEF-01", "Draft form save/submit on command kit", "Done", 100),
    ("Cursor brief", "BRIEF-02", "Table and tabs kit (DataTable, RecordTabs)", "Done", 100),
    ("Cursor brief", "BRIEF-03", "React test runner (Vitest) for kit hooks", "Done", 100),
    ("Cursor brief", "BRIEF-04", "local:check preflight, stale lock, memory gate", "Done", 100),
    ("Cursor brief", "BRIEF-05", "Shared stage screen (IAME, Reviewer, Programme, Director)", "Done", 100),
    ("Cursor brief", "BRIEF-06", "Neutral test ids (Director/Secretary approval)", "Done", 100),
    ("Cursor brief", "BRIEF-07", "Shared proposal panels (fee rules, rating schemes)", "Done", 100),
    ("Cursor brief", "BRIEF-08", "Whole-repository ESLint baseline (BL-035)", "Done", 100),
    ("Cursor brief", "BRIEF-09", "Stored state via useSyncExternalStore (BL-141)", "Pending", 0),
    ("Cursor brief", "BRIEF-10", "Notification bell and Notifications page (WP10.1b UI)", "Done", 100),
    ("Cursor brief", "BRIEF-11", "Identity strip + calmer notification refresh (BL-089)", "Pending", 0),
    ("Cursor brief", "BRIEF-12", "Business settings screen (WP08.1h frontend)", "Pending (blocked)", 0),
    ("Work package", "WP02", "Identity, session, MFA, model-application policy", "In progress", 75),
    ("Work package", "WP03", "OpenAPI contract, BFF, contract tests", "Done", 100),
    ("Work package", "WP04", "Masters and brand authorisation", "In progress", 60),
    ("Work package", "WP05.1a", "Read-only model applications UI", "Done", 100),
    ("Work package", "WP05.1b", "Draft create/edit", "Done", 100),
    ("Work package", "WP05.1c", "Draft submission", "Done", 100),
    ("Work package", "WP05.1d", "Submit-time evidence gates", "In progress", 70),
    ("Work package", "WP06.1a", "Document intake and storage", "Done", 100),
    ("Work package", "WP06.2", "Blob retention, orphans, limits (BL-001–003)", "Not started", 0),
    ("Work package", "WP07.1a", "Finance fee confirmation", "Done", 100),
    ("Work package", "WP07.1b–f", "IAME, Reviewer, Rating, Director, Secretary stages", "Done", 100),
    ("Work package", "WP07.1g–h", "Return, resubmit, rejection", "Done", 100),
    ("Work package", "WP07.1i–k", "History, inbox, workflow views", "Done", 100),
    ("Work package", "WP08.1a", "Fee rule administration", "Done", 100),
    ("Work package", "WP08.1b", "Tax line on fee rules", "Done", 100),
    ("Work package", "WP08.1c", "Rating schemes administration", "Done", 100),
    ("Work package", "WP08.1d", "Permission-based admin menu", "Done", 100),
    ("Work package", "WP08.1e", "Fee confirmation correction proposals", "Done", 100),
    ("Work package", "WP08.1f", "Fee reversal", "Done", 100),
    ("Work package", "WP08.1g", "Born-pending and notification proof", "Done", 100),
    ("Work package", "WP08.1h", "Business settings (propose/approve numbers)", "In progress", 50),
    ("Work package", "WP09.1a", "Certificate record", "Done", 100),
    ("Work package", "WP09.1b", "Printable certificate and label", "Done", 100),
    ("Work package", "WP09.1c", "Public verification page", "Done", 100),
    ("Work package", "WP09.1d", "Verification rate limit", "Done", 100),
    ("Work package", "WP10.1a", "In-portal notifications (API)", "Done", 100),
    ("Work package", "WP10.1b", "Notification bell and page", "Done", 100),
    ("Work package", "WP11", "AI demonstrations (local)", "Not started", 0),
    ("Platform", "WP01", "Documentation and design evidence", "Done", 100),
    ("Backlog theme", "BL-024/025", "Feature template and OpenAPI split", "In progress", 40),
    ("Backlog theme", "BL-040–042", "BEE owner decisions (fees, masters, evidence gates)", "Open", 0),
    ("Owner decision", "D10", "Provisional local rules for notifications (portal only)", "Accepted", 100),
]


def style_header(ws, row: int, cols: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def autosize(ws, max_width: int = 48) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col:
            if cell.value:
                length = max(length, min(max_width, len(str(cell.value))))
        ws.column_dimensions[letter].width = max(12, min(max_width, length + 2))


def main() -> None:
    if __import__("os").environ.get("BEE_REFRESH_SCREEN_CATALOG") == "1":
        export = ROOT / "scripts/tracking/export-screen-catalog.mjs"
        subprocess.run(["npx", "tsx", str(export)], cwd=str(ROOT), check=True, timeout=120)
    elif not CATALOG.is_file():
        raise SystemExit(f"Missing {CATALOG}; run: BEE_REFRESH_SCREEN_CATALOG=1 python scripts/tracking/build-portal-status-xlsx.py")
    catalog = json.loads(CATALOG.read_text())
    screens = catalog["rows"]
    apply_screen_handling(screens, load_handling_by_app_path())

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Screens"

    ws1.merge_cells("A1:G1")
    ws1["A1"] = LEGEND
    ws1["A1"].font = Font(name="Arial", italic=True, size=10)
    ws1["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws1.row_dimensions[1].height = 48

    headers1 = [
        "Screen Id",
        "Screen Name",
        "Module",
        "Applicable roles (Annex A access)",
        "Status",
        "Screen handling",
        "Remarks",
    ]
    for i, h in enumerate(headers1, start=1):
        ws1.cell(row=2, column=i, value=h)
    style_header(ws1, 2, len(headers1))

    body_font = Font(name="Arial", size=10)
    for r_idx, row in enumerate(screens, start=3):
        ws1.cell(row=r_idx, column=1, value=row["screenId"]).font = body_font
        ws1.cell(row=r_idx, column=2, value=row["screenName"]).font = body_font
        ws1.cell(row=r_idx, column=3, value=row["moduleName"]).font = body_font
        roles = row["rolesAccess"]
        if isinstance(roles, list):
            roles = ", ".join(roles) if roles else "—"
        ws1.cell(row=r_idx, column=4, value=roles).font = body_font
        ws1.cell(row=r_idx, column=5, value=row["status"]).font = body_font
        ws1.cell(row=r_idx, column=6, value=row.get("screenHandling", "")).font = body_font
        ws1.cell(row=r_idx, column=7, value=row["remarks"]).font = body_font
        for c in range(1, 8):
            ws1.cell(row=r_idx, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    ws1.freeze_panes = "A3"
    autosize(ws1)

    ws2 = wb.create_sheet("Work register")
    headers2 = ["Work Category", "Work Id", "Work Description", "Status", "% completed"]
    for i, h in enumerate(headers2, start=1):
        ws2.cell(row=1, column=i, value=h)
    style_header(ws2, 1, len(headers2))

    for r_idx, row in enumerate(TAB2_ROWS, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font = body_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws2.freeze_panes = "A2"
    autosize(ws2)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT} ({len(screens)} screens, {len(TAB2_ROWS)} work rows)")


if __name__ == "__main__":
    main()
